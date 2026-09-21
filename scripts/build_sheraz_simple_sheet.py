# -*- coding: utf-8 -*-
"""
Simple 1-sheet route list for a vendor (Sheraz) to fill in Sedan/Staria/GMC
prices and send back. No prices pre-filled - deliberately blank.

Run: python scripts/build_sheraz_simple_sheet.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\786\Documents\WEBSITES\project\taxidriver")
IN_JSON = ROOT / "data" / "sheraz-route-list.json"
OUT_XLSX = ROOT / "PRICE-BOOK-FOR-SHERAZ.xlsx"

GREEN = "16A34A"
YELLOW_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
HEADER_FILL = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SECTION_FILL = PatternFill(start_color="14532D", end_color="14532D", fill_type="solid")
SECTION_FONT = Font(color="FFFFFF", bold=True, size=13)
TITLE_FONT = Font(color="14532D", bold=True, size=16)
thin = Side(style="thin", color="DDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

CLEAN_NAME = {
    "Doha/Qatar": "Doha",
    "NEOM/Oxagon": "NEOM",
    "Aqaba/Jordan": "Aqaba",
    "Al-Ula": "AlUla",
}


def clean(label: str) -> str:
    for raw, nice in CLEAN_NAME.items():
        label = label.replace(raw, nice)
    if "(local / airport transfer)" in label:
        city = label.replace(" (local / airport transfer)", "")
        label = f"{city} City ⇄ {city} Airport"
    return label


data = json.load(open(IN_JSON, encoding="utf-8"))
cross_border = [clean(r) for r in data["crossBorder"]]
intercity = [clean(r) for r in data["intercity"]]

wb = Workbook()
ws = wb.active
ws.title = "Price Book"
ws.sheet_view.showGridLines = False

headers = ["Route", "Sedan (SAR)", "Staria (SAR)", "GMC (SAR)"]
widths = [34, 16, 16, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

# Title
ws.merge_cells("A1:D1")
ws["A1"] = "TAXI SAUDI ARABIA — PRICE BOOK (please fill in your best price per vehicle)"
ws["A1"].font = TITLE_FONT
ws.row_dimensions[1].height = 26
ws.merge_cells("A2:D2")
ws["A2"] = "Fill in Sedan / Staria / GMC prices (SAR) for each route below and send back. Leave blank if you don't offer that vehicle for a route."
ws["A2"].font = Font(italic=True, size=10, color="555555")
ws.row_dimensions[2].height = 18

row = 4


def section(title, count):
    global row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    c = ws.cell(row=row, column=1, value=f"{title}  ({count} routes)")
    c.fill = SECTION_FILL
    c.font = SECTION_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 24
    row += 1

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    row += 1


def add_rows(routes):
    global row
    for r in routes:
        ws.cell(row=row, column=1, value=r).border = BORDER
        for col in (2, 3, 4):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            cell.fill = YELLOW_FILL
        row += 1
    row += 1  # blank spacer row between sections


section("CROSS BORDER ROUTES", len(cross_border))
add_rows(cross_border)

section("INTERCITY ROUTES (CITY TO CITY)", len(intercity))
add_rows(intercity)

ws.freeze_panes = "A5"

wb.save(OUT_XLSX)
print(f"Saved {OUT_XLSX}")
print(f"Cross border rows: {len(cross_border)}, Intercity rows: {len(intercity)}")
