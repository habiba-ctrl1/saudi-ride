# -*- coding: utf-8 -*-
"""
Builds TAXI-SAUDI-PRICE-BOOK-APPROVAL.xlsx from data/pricing-book-consolidated.json
(produced by scripts/consolidate_price_book.py). Run that script first.

Run: python scripts/build_price_book_excel.py
"""
import json
import re
import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(r"C:\Users\786\Documents\WEBSITES\project\taxidriver")
IN_JSON = ROOT / "data" / "pricing-book-consolidated.json"
OUT_XLSX = ROOT / "TAXI-SAUDI-PRICE-BOOK-APPROVAL.xlsx"

GREEN = "16A34A"
YELLOW = "FACC15"
DARK = "14532D"
LIGHT_GREY = "F3F4F6"

HEADER_FILL = PatternFill(start_color=GREEN, end_color=GREEN, fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color=DARK, bold=True, size=16)
SUBTITLE_FONT = Font(color="555555", size=10, italic=True)
WARN_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
CONFLICT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
thin = Side(style="thin", color="DDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

with open(IN_JSON, "r", encoding="utf-8") as f:
    data = json.load(f)

route_families = data["routeFamilies"]
evidence = data["evidence"]
VENDOR_MARGIN = data["vendorMargin"]

wb = Workbook()
wb.remove(wb.active)


def shorten_source_ref(ref: str) -> str:
    if "price_book_entries id" in ref:
        return "TSA Price Book DB (WhatsApp chat evidence)"
    if "naimatullah" in ref.lower():
        return "Naimatullah vendor rate sheet"
    return ref


def add_title_block(ws, title, subtitle, n_cols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c2 = ws.cell(row=2, column=1, value=subtitle)
    c2.font = SUBTITLE_FONT
    ws.row_dimensions[2].height = 16


def style_header_row(ws, row_idx, headers):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row_idx].height = 30


def route_overall_status(vehicles: dict):
    statuses = [v["status"] for v in vehicles.values()]
    if "Conflicting Prices" in statuses:
        return "Conflicting Prices"
    if "Multiple Sources" in statuses:
        return "Multiple Sources"
    if statuses:
        return "Single Source"
    return "Needs Confirmation"


# ---------------------------------------------------------------------------
# SHEET 1 - APPROVAL PRICE BOOK
# ---------------------------------------------------------------------------

ws1 = wb.create_sheet("APPROVAL PRICE BOOK")
headers1 = [
    "Route", "Trip Type", "Cross Border?",
    "Sedan Lowest", "Sedan Highest",
    "GMC/SUV Lowest", "GMC/SUV Highest",
    "Staria Lowest", "Staria Highest",
    "Currency", "Pricing Status", "Approval Required", "Notes",
]
add_title_block(ws1, "TAXI SAUDI ARABIA — PRICE BOOK APPROVAL",
                 f"Consolidated from WhatsApp client-quote history + Naimatullah vendor rate card (vendor cost + SAR {VENDOR_MARGIN} margin shown as a derived reference, not a live price). Generated {datetime.date.today().isoformat()}. Nothing here is live/approved until the company confirms.",
                 len(headers1))
style_header_row(ws1, 4, headers1)

STD_CATS = ["Sedan", "GMC/SUV", "Staria"]
row_i = 5
for rf in sorted(route_families, key=lambda r: (r["routeFamily"], r["tripType"])):
    vehicles = rf["vehicles"]
    if not vehicles:
        continue
    trip_type = rf["tripType"]
    status = route_overall_status(vehicles)
    missing = [c for c in STD_CATS if c not in vehicles]
    notes_parts = []
    if missing:
        notes_parts.append("No data yet: " + ", ".join(missing))
    if rf["otherVehicleNotes"]:
        notes_parts.append("Also has " + "/".join(rf["otherVehicleNotes"]) + " quotes (see Source Details)")
    notes = "; ".join(notes_parts)

    vals = [
        rf["routeFamily"], trip_type, "Yes" if rf["crossBorder"] else "No",
        vehicles.get("Sedan", {}).get("lowest"), vehicles.get("Sedan", {}).get("highest"),
        vehicles.get("GMC/SUV", {}).get("lowest"), vehicles.get("GMC/SUV", {}).get("highest"),
        vehicles.get("Staria", {}).get("lowest"), vehicles.get("Staria", {}).get("highest"),
        "SAR", status, "Yes", notes,
    ]
    for col, v in enumerate(vals, start=1):
        cell = ws1.cell(row=row_i, column=col, value=v)
        cell.border = BORDER
        if col in (4, 5, 6, 7, 8, 9) and isinstance(v, (int, float)):
            cell.number_format = '#,##0 "SAR"'
        if col == 11:
            if v == "Conflicting Prices":
                cell.fill = CONFLICT_FILL
            elif v == "Needs Confirmation":
                cell.fill = WARN_FILL
    row_i += 1

last_row1 = row_i - 1
widths1 = [30, 12, 13, 12, 12, 13, 13, 12, 12, 10, 16, 15, 44]
for i, w in enumerate(widths1, start=1):
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.freeze_panes = "A5"
ws1.auto_filter.ref = f"A4:{get_column_letter(len(headers1))}{last_row1}"

# ---------------------------------------------------------------------------
# SHEET 2 - PRICING SOURCE DETAILS
# ---------------------------------------------------------------------------

ws2 = wb.create_sheet("PRICING SOURCE DETAILS")
headers2 = ["Route", "Pickup", "Destination", "Vehicle", "Price", "Currency", "Trip Type", "Source", "Source File / Chat Reference", "Date", "Notes"]
add_title_block(ws2, "PRICING SOURCE DETAILS", "Every individual quote/rate behind the Approval Price Book sheet — full evidence trail, nothing hidden.", len(headers2))
style_header_row(ws2, 4, headers2)

row_i = 5
for e in sorted(evidence, key=lambda x: (x["routeFamily"], x["vehicleCategory"], x["price"])):
    vals = [
        e["routeFamily"], e["pickup"], e["dropoff"], f'{e["vehicleCategory"]} ({e["vehicleRaw"]})' if e["vehicleCategory"] != e["vehicleRaw"] else e["vehicleRaw"],
        e["price"], e["currency"], "One-way" if e["tripType"] == "ONE_WAY" else "Round-trip",
        e["source"], shorten_source_ref(e["sourceRef"]), e["date"], e["notes"],
    ]
    for col, v in enumerate(vals, start=1):
        cell = ws2.cell(row=row_i, column=col, value=v)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(col == 11))
        if col == 5 and isinstance(v, (int, float)):
            cell.number_format = '#,##0'
        if e["isVendorCost"]:
            cell.fill = PatternFill(start_color=LIGHT_GREY, end_color=LIGHT_GREY, fill_type="solid")
    row_i += 1

last_row2 = row_i - 1
widths2 = [24, 26, 26, 22, 10, 10, 12, 40, 30, 12, 50]
for i, w in enumerate(widths2, start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A5"
ws2.auto_filter.ref = f"A4:{get_column_letter(len(headers2))}{last_row2}"

# ---------------------------------------------------------------------------
# SHEET 3 - ROUTE REVIEW
# ---------------------------------------------------------------------------

ws3 = wb.create_sheet("ROUTE REVIEW")
headers3 = ["Route", "Sedan (Low-High)", "GMC/SUV (Low-High)", "Staria (Low-High)", "Lowest Observed", "Highest Observed", "Number of Sources", "Conflict?", "Company Decision", "Final Approved Price", "Approval Notes"]
add_title_block(ws3, "ROUTE REVIEW — FOR COMPANY APPROVAL", "Please review each route. Fill in 'Company Decision' and 'Final Approved Price' where you agree or want to set a different rate. Leave blank where more discussion is needed.", len(headers3))
style_header_row(ws3, 4, headers3)

def route_label(rf):
    """Route family name, with trip type suffixed only when this route has both
    one-way and round-trip evidence (so the label stays clean otherwise)."""
    both = any(
        other is not rf and other["routeFamily"] == rf["routeFamily"]
        for other in route_families
    )
    return f'{rf["routeFamily"]} ({rf["tripType"]})' if both else rf["routeFamily"]


row_i = 5
for rf in sorted(route_families, key=lambda r: (r["routeFamily"], r["tripType"])):
    vehicles = rf["vehicles"]
    if not vehicles:
        continue

    def rng(cat):
        v = vehicles.get(cat)
        if not v:
            return "—"
        return f'{v["lowest"]:.0f}' if v["lowest"] == v["highest"] else f'{v["lowest"]:.0f}-{v["highest"]:.0f}'

    all_lows = [v["lowest"] for v in vehicles.values()]
    all_highs = [v["highest"] for v in vehicles.values()]
    total_sources = sum(v["sourceCount"] for v in vehicles.values())
    conflict = any(v["status"] == "Conflicting Prices" for v in vehicles.values())

    vals = [
        route_label(rf), rng("Sedan"), rng("GMC/SUV"), rng("Staria"),
        min(all_lows), max(all_highs), total_sources,
        "Yes" if conflict else "No", "", "", "",
    ]
    for col, v in enumerate(vals, start=1):
        cell = ws3.cell(row=row_i, column=col, value=v)
        cell.border = BORDER
        if col in (5, 6) and isinstance(v, (int, float)):
            cell.number_format = '#,##0 "SAR"'
        if col == 8 and v == "Yes":
            cell.fill = CONFLICT_FILL
    row_i += 1

last_row3 = row_i - 1
widths3 = [30, 18, 18, 18, 15, 15, 14, 10, 20, 18, 40]
for i, w in enumerate(widths3, start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = "A5"
ws3.auto_filter.ref = f"A4:{get_column_letter(len(headers3))}{last_row3}"

# ---------------------------------------------------------------------------
# SHEET 4 - CROSS BORDER
# ---------------------------------------------------------------------------

ws4 = wb.create_sheet("CROSS BORDER")
cb_routes = [rf for rf in route_families if rf["crossBorder"]]
has_staria = any("Staria" in rf["vehicles"] for rf in cb_routes)

headers4 = ["Route", "Country From", "Country To", "Sedan Lowest", "Sedan Highest", "GMC/SUV Lowest", "GMC/SUV Highest"]
if has_staria:
    headers4 += ["Staria Lowest", "Staria Highest"]
headers4 += ["Source Count", "Notes", "Final Approved Sedan", "Final Approved GMC/SUV"]

add_title_block(ws4, "CROSS-BORDER ROUTES", "Separated because these involve customs/border fees and driver return-trip costs — pricing here needs extra care before approval.", len(headers4))
style_header_row(ws4, 4, headers4)

row_i = 5
for rf in sorted(cb_routes, key=lambda r: (r["routeFamily"], r["tripType"])):
    vehicles = rf["vehicles"]
    total_sources = sum(v["sourceCount"] for v in vehicles.values())
    missing = [c for c in ["Sedan", "GMC/SUV"] if c not in vehicles]
    notes = ("No data yet: " + ", ".join(missing)) if missing else ""
    if any(v["status"] == "Conflicting Prices" for v in vehicles.values()):
        notes = (notes + "; " if notes else "") + "Conflicting prices observed — see Route Review"

    vals = [
        route_label(rf), rf["countryFrom"], rf["countryTo"],
        vehicles.get("Sedan", {}).get("lowest"), vehicles.get("Sedan", {}).get("highest"),
        vehicles.get("GMC/SUV", {}).get("lowest"), vehicles.get("GMC/SUV", {}).get("highest"),
    ]
    if has_staria:
        vals += [vehicles.get("Staria", {}).get("lowest"), vehicles.get("Staria", {}).get("highest")]
    vals += [total_sources, notes, "", ""]

    for col, v in enumerate(vals, start=1):
        cell = ws4.cell(row=row_i, column=col, value=v)
        cell.border = BORDER
        header = headers4[col - 1]
        if ("Lowest" in header or "Highest" in header) and isinstance(v, (int, float)):
            cell.number_format = '#,##0 "SAR"'
    row_i += 1

last_row4 = row_i - 1
widths4 = [26, 12, 12] + [13, 13, 14, 14] + ([13, 13] if has_staria else []) + [12, 40, 18, 20]
for i, w in enumerate(widths4, start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w
ws4.freeze_panes = "A5"
ws4.auto_filter.ref = f"A4:{get_column_letter(len(headers4))}{last_row4}"

for ws in (ws1, ws2, ws3, ws4):
    ws.sheet_view.showGridLines = False

wb.save(OUT_XLSX)
print(f"Saved {OUT_XLSX}")
print(f"Sheet 1 (Approval Price Book): {last_row1 - 4} routes")
print(f"Sheet 2 (Pricing Source Details): {last_row2 - 4} evidence rows")
print(f"Sheet 3 (Route Review): {last_row3 - 4} routes")
print(f"Sheet 4 (Cross Border): {last_row4 - 4} routes, Staria column {'included' if has_staria else 'omitted (no evidence)'}")
